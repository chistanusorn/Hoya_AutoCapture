import sys
import json
from docx import Document
from docx.oxml.ns import qn

class NumberingTracker:
    def __init__(self, doc):
        self.nums = {} 
        self.abstract_nums = {} 
        self.state = {}
        
        self._parse_numbering(doc)
        self.p_to_num = {} 
        self._traverse_doc(doc)
        
    def _parse_numbering(self, doc):
        try:
            numbering_part = doc.part.numbering_part
        except NotImplementedError:
            numbering_part = None
            
        if not numbering_part:
            return
            
        element = numbering_part.element
        for num in element.findall(qn('w:num')):
            numId = num.get(qn('w:numId'))
            absNum = num.find(qn('w:abstractNumId'))
            if absNum is not None:
                self.nums[numId] = absNum.get(qn('w:val'))
                
        for abs_num in element.findall(qn('w:abstractNum')):
            absNumId = abs_num.get(qn('w:abstractNumId'))
            levels = {}
            for lvl in abs_num.findall(qn('w:lvl')):
                ilvl = lvl.get(qn('w:ilvl'))
                start = lvl.find(qn('w:start'))
                start_val = int(start.get(qn('w:val'))) if start is not None else 1
                
                fmt = lvl.find(qn('w:numFmt'))
                fmt_val = fmt.get(qn('w:val')) if fmt is not None else 'decimal'
                
                text = lvl.find(qn('w:lvlText'))
                text_val = text.get(qn('w:val')) if text is not None else ''
                
                levels[ilvl] = {'start': start_val, 'fmt': fmt_val, 'text': text_val}
            self.abstract_nums[absNumId] = levels
            self.state[absNumId] = {}

    def _traverse_doc(self, doc):
        for p in doc.element.xpath('//w:p'):
            numPr = p.find(qn('w:pPr/w:numPr'))
            if numPr is None:
                continue
                
            numId_el = numPr.find(qn('w:numId'))
            ilvl_el = numPr.find(qn('w:ilvl'))
            
            if numId_el is None:
                continue
                
            numId = numId_el.get(qn('w:val'))
            ilvl = ilvl_el.get(qn('w:val')) if ilvl_el is not None else '0'
            
            if numId not in self.nums:
                continue
                
            absNumId = self.nums[numId]
            if absNumId not in self.abstract_nums or ilvl not in self.abstract_nums[absNumId]:
                continue
                
            lvl_info = self.abstract_nums[absNumId][ilvl]
            
            st = self.state[absNumId]
            ilvl_int = int(ilvl)
            
            if ilvl not in st:
                st[ilvl] = lvl_info['start']
            else:
                st[ilvl] += 1
                
            for k in list(st.keys()):
                if int(k) > ilvl_int:
                    del st[k]
                    
            if lvl_info['fmt'] == 'decimal':
                txt = lvl_info['text']
                for i in range(ilvl_int + 1):
                    lvl_str = str(i)
                    val = st.get(lvl_str, self.abstract_nums[absNumId].get(lvl_str, {}).get('start', 1))
                    txt = txt.replace(f'%{i+1}', str(val))
                
                # lxml element can be used as key? Yes, in python dictionaries.
                self.p_to_num[p] = txt

    def get_numbering(self, paragraph):
        return self.p_to_num.get(paragraph._p, "")


def test(docx_path):
    doc = Document(docx_path)
    tracker = NumberingTracker(doc)
    
    for i, table in enumerate(doc.tables):
        print(f"Table {i}:")
        for ri, row in enumerate(table.rows):
            cell = row.cells[0]
            text = cell.text.strip()
            
            # test numbering logic
            if not text:
                for p in cell.paragraphs:
                    num = tracker.get_numbering(p)
                    if num:
                        text = num
                        break
            
            if text:
                print(f"  Row {ri}: {text}")

if __name__ == '__main__':
    test(sys.argv[1])
