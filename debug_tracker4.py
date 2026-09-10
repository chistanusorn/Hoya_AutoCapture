import sys
from docx import Document
from word_target import NumberingTracker
from docx.oxml.ns import qn

def debug():
    doc_path = r"d:\hoya\AutoCapture\output\VPR-2603-003_F-CP.GE-X34.01_Validation Report_PQ_FogGuard_draft_V1_20260324_capture.docx"
    doc = Document(doc_path)
    tracker = NumberingTracker(doc)
    
    table = doc.tables[10]
    ns_w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    
    for ri, row in enumerate(table.rows):
        if not row.cells: continue
        c0 = row.cells[0]
        
        for p in c0.paragraphs:
            numPr = p._p.find(f'.//{ns_w}numPr')
            if numPr is not None:
                numId_el = numPr.find(f'.//{ns_w}numId')
                ilvl_el = numPr.find(f'.//{ns_w}ilvl')
                
                numId = numId_el.get(f'{ns_w}val') if numId_el is not None else None
                ilvl = ilvl_el.get(f'{ns_w}val') if ilvl_el is not None else '0'
                
                print(f"Row {ri} numPr: numId={numId}, ilvl={ilvl}")
                
                if numId in tracker.nums:
                    absNumId = tracker.nums[numId]
                    print(f"  absNumId={absNumId}")
                    if absNumId in tracker.abstract_nums:
                        lvl = tracker.abstract_nums[absNumId].get(ilvl)
                        print(f"  lvl_info={lvl}")
            else:
                print(f"Row {ri} NO numPr")
                
            print(f"  Calculated by tracker: '{tracker.get_numbering(p)}'")

if __name__ == "__main__":
    debug()
