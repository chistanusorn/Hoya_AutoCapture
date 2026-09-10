import sys
from docx import Document
from word_target import NumberingTracker, actual_column

def debug():
    doc = Document(r"d:\hoya\AutoCapture\output\VPR-2603-003_F-CP.GE-X34.01_Validation Report_PQ_FogGuard_draft_V1_20260324_capture.docx")
    tracker = NumberingTracker(doc)
    
    for ti, table in enumerate(doc.tables):
        if actual_column(table) is None:
            continue
            
        print(f"--- Table {ti} ---")
        for ri, row in enumerate(table.rows):
            if not row.cells: continue
            
            c0 = row.cells[0]
            c1 = row.cells[1] if len(row.cells) > 1 else None
            
            nums0 = [tracker.get_numbering(p) for p in c0.paragraphs]
            nums1 = [tracker.get_numbering(p) for p in c1.paragraphs] if c1 else []
            
            text0 = c0.text.strip().replace('\n', ' ')
            text1 = c1.text.strip().replace('\n', ' ') if c1 else ""
            
            print(f"Row {ri}:")
            print(f"  Col 0 (No): text='{text0}', nums={nums0}")
            print(f"  Col 1 (Desc): text='{text1}', nums={nums1}")
            
if __name__ == "__main__":
    debug()
