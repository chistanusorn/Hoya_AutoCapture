"""โหมด <pic> — วางรูปแทน marker <pic> ในตาราง

ใช้กับเอกสารที่คนทำ format ไว้: ใส่ <pic> ตรงที่ต้องการรูป
ชื่อของแต่ละจุดใช้เลขจากคอลัมน์ No. ของแถวนั้น ถ้าแถวเดียวมีหลาย <pic>
เติมลำดับย่อย เช่น 1.1, 1.1 (2), 1.1 (3)

ยืดหยุ่นกว่าโหมดขั้นตอน เพราะไม่ต้องเดาคอลัมน์/ตัวหนา/รูปแบบเลข
แค่หา <pic> แล้ววางทับ
"""

import re
import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

import config
from word_target import NumberingTracker, _is_step_number, place_full_width

PIC = "<pic>"


def has_pic_markers(path):
    """เอกสารนี้ใช้โหมด <pic> ไหม — มี <pic> ในตารางอย่างน้อย 1 อัน"""
    doc = Document(str(path))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if PIC in cell.text.lower():
                    return True
    return False


def _table_label(recent_lines):
    """ป้าย dropdown — เอา 'Table N ...' ถ้าเจอ ไม่งั้นบรรทัดล่าสุด"""
    for line in reversed(recent_lines):
        if re.match(r"table\s*\d+", line, re.IGNORECASE):
            return line[:40]
    return (recent_lines[-1][:40] if recent_lines else "")


def find_pic_tables(template):
    """ตารางที่มี <pic> — คืน [(table_index, ป้ายชื่อ, จำนวน pic)]

    เฉพาะ <pic> ในตาราง — นอกตารางไม่สนใจ (เป็นรูปประกอบของคนเขียน)
    """
    doc = Document(str(template))
    recent = []
    ti = 0
    labels = {}
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            t = "".join(n.text or "" for n in child.iter(qn("w:t"))).strip()
            if t:
                recent.append(t)
                if len(recent) > 6:
                    recent.pop(0)
        elif child.tag == qn("w:tbl"):
            labels[ti] = _table_label(recent)
            ti += 1

    out = []
    tracker = NumberingTracker(doc)
    for i, table in enumerate(doc.tables):
        # นับ slot (รวมที่วางรูปแล้ว) ไม่ใช่แค่ <pic> ที่เหลือ
        # ไม่งั้นตารางที่ใส่รูปครบจะหายจาก dropdown ตอนเปิดใหม่
        n = len(_pic_slots(table, tracker))
        if not n:
            continue
        # ข้ามตาราง metadata (เช่นช่อง No. เป็น 'VALIDATION TYPE:') ที่บังเอิญ
        # มี <pic> — เอาเฉพาะตารางผลที่ No. เป็นเลขขั้นตอนอย่างน้อย 1 แถว
        if not _has_step_row(table, tracker):
            continue
        out.append((i, labels.get(i) or f"ตารางที่ {i}", n))
    return out


def _has_step_row(table, tracker):
    """ตารางนี้มีแถวที่ช่อง No. เป็นเลขขั้นตอน (1, 1.1) อย่างน้อย 1 ไหม"""
    for row in table.rows:
        no = tracker.cell_number(row.cells[0]) if tracker else row.cells[0].text.strip().rstrip(".")
        if _is_step_number(no):
            return True
    return False


def _caption(cell, para_idx):
    """ข้อความบรรทัดสุดท้ายเหนือรูป — บอกว่ารูปนี้คือรูปอะไร

    ในเอกสารจริงคนเขียนอธิบายไว้เหนือรูปเสมอ (เช่น 'Input username: rx3fogA
    and password') จึงเจาะจงกว่าคอลัมน์ Description ที่ใช้ร่วมกันทั้งแถว
    รูปที่วางติดกันโดยไม่มีข้อความคั่นจะได้ข้อความเดียวกัน — แยกด้วยลำดับ (2) (3)
    """
    for para in reversed(cell.paragraphs[:para_idx]):
        text = para.text.strip()
        if text and text.lower() != PIC:
            return text
    return None


def _desc_column(table):
    """คอลัมน์ Description/Action จากแถวหัวตาราง — None ถ้าไม่มี"""
    if not table.rows:
        return None
    for i, cell in enumerate(table.rows[0].cells):
        head = cell.text.strip().lower()
        if "description" in head or "action" in head:
            return i
    return None


def _pic_slots(table, tracker=None):
    """ทุก 'จุดรูป' ในตาราง เรียงตามลำดับที่เจอ — ทั้งที่ยังเป็น <pic>
    และที่วางรูปไปแล้ว เพื่อให้ลำดับ (slot_no) คงที่ ไม่เลื่อนหลังวางรูป

    ชื่อ = เลข No. ของแถว ถ้าแถวมีหลาย <pic> เติมลำดับย่อย: 1.1, 1.1 (2)...

    คืน [(cell, ดัชนี paragraph, ชื่อ, มีรูปแล้วไหม, ข้อความกำกับรูป)]
    """
    desc_col = _desc_column(table)
    slots = []
    for row in table.rows:
        no = tracker.cell_number(row.cells[0]) if tracker else row.cells[0].text.strip().rstrip(".")

        # หา <pic> / รูป ในแถวนี้ก่อน เพื่อรู้ว่าต้องเติมลำดับย่อยไหม
        found = []
        for cell in row.cells:
            for pi, para in enumerate(cell.paragraphs):
                if "graphicData" in para._p.xml:
                    found.append((cell, pi, True))
                elif para.text.strip().lower() == PIC:
                    found.append((cell, pi, False))

        # เอกสารที่ไม่ได้เขียนอธิบายเหนือรูป ยังพอมีคอลัมน์ Description ให้ใช้แทน
        row_desc = None
        if desc_col is not None and desc_col < len(row.cells):
            row_desc = row.cells[desc_col].text.strip().replace("\n", " ") or None

        multi = len(found) > 1
        for n, (cell, pi, placed) in enumerate(found, start=1):
            # อันแรกของแถวใช้เลขเปล่า อันถัดไปเติม (2) (3)...
            label = f"{no} ({n})" if multi and n > 1 else no
            slots.append((cell, pi, label, placed, _caption(cell, pi) or row_desc))
    return slots


def load_slots(template, table_index):
    """รายการจุดรูปในตารางที่เลือก

    คืน [(ลำดับ, ชื่อ, จำนวนรวม, มีรูปแล้วไหม, ข้อความกำกับรูป)]
    """
    ensure_working_copy(template)
    doc = Document(str(config.word_working(template)))
    table = doc.tables[table_index]
    tracker = NumberingTracker(doc)
    slots = _pic_slots(table, tracker)
    total = len(slots)
    return [(i + 1, name or f"รูปที่ {i + 1}", total, placed, caption)
            for i, (_, _, name, placed, caption) in enumerate(slots)]


def ensure_working_copy(template):
    template = Path(template)
    if not template.exists():
        raise FileNotFoundError(f"ไม่พบไฟล์ต้นฉบับ: {template}")
    working = config.word_working(template)
    if not working.exists():
        shutil.copy2(template, working)
    return working


def place_image(template, table_index, slot_no, image_path):
    """วางรูปแทน <pic> ลำดับที่ slot_no (นับจาก 1) — คืน (สำเร็จ, ข้อความ)"""
    path = ensure_working_copy(template)
    doc = Document(str(path))
    table = doc.tables[table_index]
    tracker = NumberingTracker(doc)
    slots = _pic_slots(table, tracker)

    if not (1 <= slot_no <= len(slots)):
        return False, f"ไม่พบจุดรูปลำดับที่ {slot_no}"

    cell, para_idx, name, _placed, _caption = slots[slot_no - 1]
    para = cell.paragraphs[para_idx]
    for run in list(para.runs):
        run.text = ""
    place_full_width(para, image_path, table, cell)

    doc.save(str(path))
    return True, name or f"รูปที่ {slot_no}"


def slot_images(template, table_index):
    """คืน {slot_no: bytes ของรูปที่ฝังอยู่} ของตารางนี้

    ใช้เทียบว่ารูปไหนในโฟลเดอร์ถูกใส่ไปแล้ว — อ่านจากเอกสารตรงๆ จึงรู้ย้อนหลังได้
    แม้รูปจะถูกใส่ไว้ตั้งแต่ก่อนมีฟีเจอร์นี้ หรือมีคนไปแก้เอกสารเอง
    """
    working = config.word_working(template)
    if not working.exists():
        return {}

    doc = Document(str(working))
    table = doc.tables[table_index]
    tracker = NumberingTracker(doc)
    parts = doc.part.related_parts

    out = {}
    for i, (cell, para_idx, _name, placed, _cap) in enumerate(_pic_slots(table, tracker),
                                                              start=1):
        if not placed:
            continue
        para = cell.paragraphs[para_idx]
        for blip in para._p.iter(qn("a:blip")):
            rid = blip.get(qn("r:embed"))
            if rid and rid in parts:
                out[i] = parts[rid].blob
                break
    return out


def revert_image(template, table_index, slot_no):
    """เอารูปออกจาก slot ที่ระบุ คืนกลับเป็น <pic> เหมือนก่อนวาง

    ไม่ลบไฟล์ .png ใน captures\\ — ถอดออกจาก Word เท่านั้น คืน (สำเร็จ, ข้อความ)
    """
    path = ensure_working_copy(template)
    doc = Document(str(path))
    table = doc.tables[table_index]
    tracker = NumberingTracker(doc)
    slots = _pic_slots(table, tracker)

    if not (1 <= slot_no <= len(slots)):
        return False, f"ไม่พบจุดรูปลำดับที่ {slot_no}"

    cell, para_idx, name, placed, _caption = slots[slot_no - 1]
    if not placed:
        return False, "จุดนี้ยังไม่มีรูป"

    para = cell.paragraphs[para_idx]
    for run in list(para.runs):
        for drawing in run._r.findall(qn("w:drawing")):
            run._r.remove(drawing)
        run.text = ""
    para.paragraph_format.left_indent = None
    if para.runs:
        para.runs[0].text = PIC
    else:
        para.add_run(PIC)

    doc.save(str(path))
    return True, name or f"รูปที่ {slot_no}"


